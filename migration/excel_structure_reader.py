from openpyxl import load_workbook


def get_table_names(excel_file):

    workbook = load_workbook(
        excel_file,
        read_only=True
    )

    worksheet = workbook["Tables_Fields"]

    table_names = set()

    for row in worksheet.iter_rows(min_row=2, values_only=True):

        table_name = row[0]

        if table_name:
            table_names.add(table_name)

    workbook.close()

    return list(table_names)


def get_table_fields(excel_file, table_name):

    workbook = load_workbook(
        excel_file,
        read_only=True
    )

    worksheet = workbook["Tables_Fields"]

    fields = []

    for row in worksheet.iter_rows(min_row=2, values_only=True):

        current_table_name = row[0]

        if current_table_name == table_name:

            field = {
                "name": row[1],
                "data_type": row[2],
                "field_size": row[3],
                "required": row[4],
                "allow_zero_length": row[5],
                "default_value": row[6],
                "description": row[7],
                "primary_key": row[8],
                "auto_increment": row[9]
            }

            fields.append(field)

    workbook.close()

    return fields



def get_relationships(excel_file):
    workbook = load_workbook(
        excel_file,
        read_only=True
    )

    worksheet = workbook["Relationships"]

    relationships = []

    for row in worksheet.iter_rows(
        min_row=2,
        values_only=True
    ):
        relation_name = row[0]
        parent_table = row[1]
        parent_field = row[2]
        child_table = row[3]
        child_field = row[4]
        attributes = row[5]

        if parent_table and child_table:

            relationship = {
                "relation_name": relation_name,
                "parent_table": parent_table,
                "parent_field": parent_field,
                "child_table": child_table,
                "child_field": child_field,
                "attributes": attributes
            }

            relationships.append(
                relationship
            )

    workbook.close()

    return relationships



def get_child_tables(relationships, table_name):
    child_tables = []

    for relationship in relationships:
        if relationship["parent_table"] == table_name:
            child_table = relationship["child_table"]

            if child_table not in child_tables:
                child_tables.append(child_table)

    return child_tables



def get_all_dependent_tables(relationships, table_name):
    dependent_tables = []

    def find_children(current_table):
        child_tables = get_child_tables(
            relationships,
            current_table
        )

        for child_table in child_tables:

            if child_table not in dependent_tables:

                dependent_tables.append(child_table)

                find_children(child_table)

    find_children(table_name)

    return dependent_tables



def get_delete_order(relationships, table_name):
    dependent_tables = get_all_dependent_tables(
        relationships,
        table_name
    )

    dependent_tables.reverse()

    delete_order = dependent_tables + [table_name]

    return delete_order



def get_child_relationships(relationships, table_name):
    child_relationships = []

    for relationship in relationships:

        if relationship["parent_table"] == table_name:

            child_relationships.append(relationship)

    return child_relationships



def get_parent_relationship(relationships, table_name):
    parent_relationships = []

    for relationship in relationships:

        if relationship["child_table"] == table_name:

            parent_relationships.append(relationship)

    return parent_relationships



def get_primary_key_fields(fields):
    primary_key_fields = []

    for field in fields:

        if field["primary_key"]:
            primary_key_fields.append(
                field["name"]
            )

    return primary_key_fields



def get_all_primary_keys(excel_file):
    table_names = get_table_names(excel_file)

    primary_keys = {}

    for table_name in table_names:

        fields = get_table_fields(
            excel_file,
            table_name
        )

        primary_key_fields = get_primary_key_fields(
            fields
        )

        if primary_key_fields:
            primary_keys[table_name] = primary_key_fields

    return primary_keys



def get_indexes(excel_file):

    workbook = load_workbook(
        excel_file,
        read_only=True
    )

    worksheet = workbook["Indexes"]

    indexes = []

    for row in worksheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        table_name = row[0]
        index_name = row[1]
        primary = row[2]
        unique = row[3]
        ignore_nulls = row[4]
        field_name = row[5]

        if table_name and index_name and field_name:

            index = {
                "table_name": table_name,
                "index_name": index_name,
                "primary": primary,
                "unique": unique,
                "ignore_nulls": ignore_nulls,
                "field_name": field_name
            }

            indexes.append(
                index
            )

    workbook.close()

    return indexes



def get_table_indexes(
    excel_file,
    table_name
):

    indexes = get_indexes(
        excel_file
    )

    table_indexes = []

    for index in indexes:

        if index["table_name"] == table_name:

            table_indexes.append(
                index
            )

    return table_indexes