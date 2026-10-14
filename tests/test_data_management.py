"""0.9.0 backup/restore/export tests: all destructive operations use temp roots."""

import os
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

import hashlib
import json
import sqlite3
import stat
import tempfile
import unittest
import zipfile
from contextlib import closing
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication, QFileDialog, QMainWindow, QMessageBox, QPushButton

from database.initializer import create_empty_database
from database.runtime import runtime_connection
from gui.data_management import AccessMigrationDialog, DataManagementController, run_operation
from gui.settings_page import SettingsPage
from services.backup import BackupValidationError, DATABASE_ENTRY, SETTINGS_ENTRY, RestoreError, RestoreRecoveryRequiredError
from services.data_management import DataManagementService
from services.settings import SettingsService
from version import APP_NAME, APP_VERSION, DB_SCHEMA_VERSION, BACKUP_FORMAT_VERSION


class DataManagementTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        directory = tempfile.TemporaryDirectory(prefix='TenderSystem_data_management_')
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.db = self.root / 'current.db'
        create_empty_database(self.db)
        with runtime_connection(self.db, write=True) as connection:
            connection.executemany('INSERT INTO 人员表(人员ID,姓名,身份证号,联系电话,学历,毕业院校,专业,进入系统日期) '
                                   'VALUES(?,?,?,?,?,?,?,?)', [
                (11, '张三', '001234567890123456', '00123456789', '本科', '甲大学', '电力', '2024-01-02'),
                (20, '=literal name', None, '', None, None, None, None),
            ])
            connection.execute("INSERT INTO 人员证书表(证书ID,人员ID,证书名称) VALUES(9,11,'不应导出的证书')")
        self.images = self.root / 'Images'
        self.images.mkdir()
        (self.images / 'legacy.jpg').write_bytes(b'legacy image')
        (self.images / 'sub').mkdir()
        (self.images / 'sub' / 'page.jpg').write_bytes(b'nested page')
        (self.images / 'empty').mkdir()
        self.settings_path = self.root / SETTINGS_ENTRY
        self.settings_path.parent.mkdir()
        self.settings_path.write_text('{"dpi":300,"future":"保留"}', encoding='utf-8')
        self.service = DataManagementService(self.db, data_root=self.root)
        self.backups = self.service.backups

    def state(self):
        return (self.db.read_bytes(),
                {path.relative_to(self.images).as_posix(): path.read_bytes() for path in self.images.rglob('*') if path.is_file()},
                self.settings_path.read_bytes() if self.settings_path.exists() else None)

    def create(self, name='backup.zip'):
        return self.backups.create_backup(self.root / name)

    def rewrite(self, source, *, manifest_changes=None, remove=(), extras=(), database_edit=None):
        target = self.root / 'modified.zip'
        with zipfile.ZipFile(source) as archive:
            payload = {info.filename: archive.read(info) for info in archive.infolist()}
        manifest = json.loads(payload['backup.json'])
        if manifest_changes:
            manifest.update(manifest_changes)
        if database_edit:
            path = self.root / 'edited-snapshot.db'
            path.write_bytes(payload[DATABASE_ENTRY])
            with closing(sqlite3.connect(path)) as connection:
                with connection:
                    database_edit(connection)
            payload[DATABASE_ENTRY] = path.read_bytes()
            manifest.pop('files', None)
        payload['backup.json'] = json.dumps(manifest).encode('utf-8')
        for name in remove:
            payload.pop(name, None)
        with zipfile.ZipFile(target, 'w') as archive:
            for name, data in payload.items():
                archive.writestr(name, data)
            for name, data in extras:
                archive.writestr(name, data)
        return target

    def test_versions_are_independent_and_runtime_stays_v1(self):
        self.assertEqual('Nekozzz', APP_NAME)
        self.assertEqual('0.9.4', APP_VERSION)
        self.assertEqual(1, DB_SCHEMA_VERSION)
        self.assertEqual(1, BACKUP_FORMAT_VERSION)
        with runtime_connection(self.db) as connection:
            self.assertEqual(DB_SCHEMA_VERSION, connection.execute('PRAGMA user_version').fetchone()[0])

    def test_backup_manifest_snapshot_images_settings_and_source_unchanged(self):
        before = self.state()
        path = self.create()
        info = self.backups.inspect_backup(path)
        self.assertEqual(APP_VERSION, info.manifest['app_version'])
        self.assertEqual(BACKUP_FORMAT_VERSION, info.manifest['format_version'])
        self.assertEqual(DB_SCHEMA_VERSION, info.manifest['database_version'])
        self.assertTrue(info.manifest['includes_settings'])
        self.assertTrue(info.manifest['includes_images'])
        from datetime import datetime
        self.assertIsNotNone(datetime.fromisoformat(info.manifest['created_at']).tzinfo)
        with zipfile.ZipFile(path) as archive:
            self.assertTrue({'backup.json', DATABASE_ENTRY, SETTINGS_ENTRY, 'Images/',
                             'Images/legacy.jpg', 'Images/sub/page.jpg', 'Images/empty/'}.issubset(archive.namelist()))
            snapshot = self.root / 'check.db'
            snapshot.write_bytes(archive.read(DATABASE_ENTRY))
        with runtime_connection(snapshot) as connection:
            self.assertEqual(2, connection.execute('SELECT COUNT(*) FROM 人员表').fetchone()[0])
        self.assertEqual(before, self.state())
        self.assertFalse(list(self.root.glob('.*.tmp')))

    def test_backup_without_settings_or_images_directory(self):
        self.settings_path.unlink()
        import shutil
        shutil.rmtree(self.images)
        info = self.backups.inspect_backup(self.create())
        self.assertFalse(info.manifest['includes_settings'])
        self.assertTrue(info.manifest['includes_images'])

    def test_backup_reads_committed_wal_without_copying_main_file(self):
        connection = sqlite3.connect(self.db)
        self.addCleanup(connection.close)
        connection.execute('PRAGMA journal_mode=WAL')
        connection.execute('PRAGMA wal_autocheckpoint=0')
        connection.execute("INSERT INTO 人员表(人员ID,姓名) VALUES(99,'WAL人员')")
        connection.commit()
        before = hashlib.sha256(self.db.read_bytes()).digest()
        path = self.create()
        snapshot = self.root / 'wal-snapshot.db'
        with zipfile.ZipFile(path) as archive:
            snapshot.write_bytes(archive.read(DATABASE_ENTRY))
        with runtime_connection(snapshot) as snap:
            self.assertEqual('WAL人员', snap.execute('SELECT 姓名 FROM 人员表 WHERE 人员ID=99').fetchone()[0])
            self.assertEqual('delete', snap.execute('PRAGMA journal_mode').fetchone()[0])
        self.assertEqual(before, hashlib.sha256(self.db.read_bytes()).digest())

    def test_existing_backup_is_refused_and_failure_not_published(self):
        path = self.create()
        before = path.read_bytes()
        with self.assertRaises(FileExistsError):
            self.backups.create_backup(path)
        self.assertEqual(before, path.read_bytes())
        with self.assertRaises(ValueError):
            self.backups.create_backup(self.images / 'recursive.zip')
        with patch('services.backup._validate_database', side_effect=ValueError('injected')):
            with self.assertRaises(ValueError):
                self.create('failed.zip')
        self.assertFalse((self.root / 'failed.zip').exists())

    def test_bad_versions_names_and_manifest_types_rejected(self):
        source = self.create()
        for changes in ({'format_version': 9}, {'database_version': 9}, {'format_version': True},
                        {'app_name': 'Other'}, {'created_at': 'bad'}, {'includes_settings': 'yes'},
                        {'database_file': '../current.db'}, {'includes_images': False}):
            with self.subTest(changes=changes), self.assertRaises(BackupValidationError):
                self.backups.inspect_backup(self.rewrite(source, manifest_changes=changes))

    def test_missing_database_manifest_settings_images_rejected(self):
        source = self.create()
        for names in ((DATABASE_ENTRY,), ('backup.json',), (SETTINGS_ENTRY,),
                      ('Images/', 'Images/legacy.jpg', 'Images/sub/page.jpg', 'Images/empty/')):
            with self.subTest(names=names), self.assertRaises(BackupValidationError):
                self.backups.inspect_backup(self.rewrite(source, remove=names))

    def test_invalid_json_zip_and_sqlite_rejected(self):
        source = self.create()
        for name, content in (('bad.zip', b'not a zip'),):
            path = self.root / name
            path.write_bytes(content)
            with self.assertRaises(BackupValidationError):
                self.backups.inspect_backup(path)
        path = self.root / 'bad-json.zip'
        with zipfile.ZipFile(path, 'w') as archive:
            archive.writestr('Images/', '')
            archive.writestr('backup.json', b'{bad')
        with self.assertRaises(BackupValidationError):
            self.backups.inspect_backup(path)
        path = self.rewrite(source, database_edit=lambda c: c.execute('PRAGMA user_version=99'))
        with self.assertRaisesRegex(BackupValidationError, 'user_version'):
            self.backups.inspect_backup(path)
        path = self.rewrite(source, database_edit=lambda c: c.execute('DROP TABLE 人员证书图表'))
        with self.assertRaises(BackupValidationError):
            self.backups.inspect_backup(path)

    def test_traversal_absolute_ads_duplicate_and_symlink_rejected(self):
        source = self.create()
        for name in ('../escape.txt', '/escape.txt', 'C:/escape.txt', 'Images/../../escape.txt',
                     'Images\\escape.jpg', 'Images/page.jpg:stream', 'Images/CON', 'Images/../x',
                     'Images/LEGACY.JPG'):
            with self.subTest(name=name), self.assertRaises(BackupValidationError):
                self.backups.inspect_backup(self.rewrite(source, extras=[(name, b'bad')]))
        symlink = zipfile.ZipInfo('Images/link')
        symlink.create_system = 3
        symlink.external_attr = (stat.S_IFLNK | 0o777) << 16
        with self.assertRaises(BackupValidationError):
            self.backups.inspect_backup(self.rewrite(source, extras=[(symlink, b'../../outside')]))
        self.assertFalse((self.root.parent / 'escape.txt').exists())

    def test_payload_hash_mismatch_rejected(self):
        source = self.create()
        info = self.backups.inspect_backup(source)
        files = dict(info.manifest['files'])
        files['Images/legacy.jpg'] = '0' * 64
        with self.assertRaisesRegex(BackupValidationError, '校验失败'):
            self.backups.inspect_backup(self.rewrite(source, manifest_changes={'files': files}))

    def change_live_data(self):
        with runtime_connection(self.db, write=True) as connection:
            connection.execute("UPDATE 人员表 SET 姓名='后来修改' WHERE 人员ID=11")
        (self.images / 'legacy.jpg').write_bytes(b'new image')
        (self.images / 'extra.jpg').write_bytes(b'should not remain')
        self.settings_path.write_text('{"dpi":150}', encoding='utf-8')

    def test_restore_replaces_images_settings_database_and_creates_safety(self):
        old = self.state()
        archive = self.create()
        with zipfile.ZipFile(archive) as saved:
            expected_database = saved.read(DATABASE_ENTRY)
        self.change_live_data()
        live = self.state()
        result = self.backups.restore_backup(archive)
        self.assertEqual(expected_database, self.state()[0])
        self.assertEqual(old[1:], self.state()[1:])
        self.assertTrue((self.images / 'empty').is_dir())
        self.assertEqual(self.root / 'Backups', result.safety_backup.parent)
        self.assertTrue(result.settings_restored)
        safety = self.backups.inspect_backup(result.safety_backup)
        self.assertEqual(APP_VERSION, safety.manifest['app_version'])
        with zipfile.ZipFile(result.safety_backup) as saved:
            self.assertEqual(live[2], saved.read(SETTINGS_ENTRY))
            self.assertEqual(b'should not remain', saved.read('Images/extra.jpg'))
        self.assertFalse(list(self.root.rglob('.TenderSystem-restore-*')))

    def test_restore_without_settings_preserves_current_settings(self):
        self.settings_path.unlink()
        archive = self.create()
        self.settings_path.write_text('{"dpi":150}', encoding='utf-8')
        self.change_live_data()
        before = self.settings_path.read_bytes()
        result = self.backups.restore_backup(archive)
        self.assertFalse(result.settings_restored)
        self.assertEqual(before, self.settings_path.read_bytes())
        self.assertFalse((self.images / 'extra.jpg').exists())

    def test_safety_failure_and_invalid_backup_leave_live_data_unchanged(self):
        archive = self.create()
        self.change_live_data()
        before = self.state()
        with patch.object(self.backups, 'create_backup', side_effect=OSError('safety failure')):
            with self.assertRaises(OSError):
                self.backups.restore_backup(archive)
        self.assertEqual(before, self.state())
        invalid = self.rewrite(archive, manifest_changes={'format_version': 999})
        with self.assertRaises(BackupValidationError):
            self.backups.restore_backup(invalid)
        self.assertEqual(before, self.state())

    def test_mid_replace_failure_rolls_back_all_live_components(self):
        archive = self.create()
        self.change_live_data()
        before = self.state()
        original = os.replace
        def fail_new_images(source, target):
            if Path(target) == self.images and Path(source).name == 'new':
                raise OSError('injected Images publish failure')
            return original(source, target)
        with patch('services.backup.os.replace', side_effect=fail_new_images):
            with self.assertRaisesRegex(RestoreError, '原文件已回退'):
                self.backups.restore_backup(archive)
        self.assertEqual(before, self.state())
        self.assertEqual(1, len(list((self.root / 'Backups').glob('*.zip'))))

    def test_preview_changed_zip_rejected_before_safety(self):
        source = self.create()
        info = self.backups.inspect_backup(source)
        source.write_bytes(b'changed')
        before = self.state()
        with self.assertRaisesRegex(BackupValidationError, '预览后发生变化'):
            self.backups.restore_backup(source, expected_digest=info.sha256)
        self.assertEqual(before, self.state())
        self.assertFalse((self.root / 'Backups').exists())

    def test_restore_retires_stale_sidecars_without_applying_them(self):
        archive = self.create()
        sidecar = Path(str(self.db) + '-shm')
        sidecar.write_bytes(b'old shared memory')
        self.backups.restore_backup(archive)
        self.assertFalse(sidecar.exists())

    def test_personnel_xlsx_real_fields_text_and_no_certificates(self):
        before = self.state()
        target, count = self.service.export_people(self.root / 'people.xlsx')
        self.assertEqual(2, count)
        from contextlib import closing
        with closing(load_workbook(target)) as workbook:
            sheet = workbook.active
            self.assertEqual(3, sheet.max_row)
            self.assertEqual(8, sheet.max_column)
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual('人员ID（系统字段，请勿修改）', headers[0])
            self.assertEqual(11, sheet.cell(2, 1).value)
            self.assertEqual('001234567890123456', sheet.cell(2, 3).value)
            self.assertEqual('s', sheet.cell(2, 3).data_type)
            self.assertEqual('@', sheet.cell(2, 3).number_format)
            self.assertEqual('2024-01-02', sheet.cell(2, 8).value)
            self.assertEqual('=literal name', sheet.cell(3, 2).value)
            self.assertEqual('s', sheet.cell(3, 2).data_type)
            self.assertEqual('A2', sheet.freeze_panes)
            self.assertEqual('A1:H3', sheet.auto_filter.ref)
            self.assertFalse(any('证书' in header for header in headers))
            self.assertEqual(['人员表'], workbook.sheetnames)
        self.assertEqual(before, self.state())
        with self.assertRaises(FileExistsError):
            self.service.export_people(target)

    def test_access_service_delegates_and_never_replaces_current_or_existing(self):
        source = self.root / 'source.accdb'
        source.write_bytes(b'mocked Access source')
        before = self.state()
        original = source.read_bytes()
        target = self.root / 'new.db'
        def fake_migration(source_path, target_path, *, excel_file):
            self.assertEqual(self.service.migration_resource, excel_file)
            create_empty_database(target_path)
            return {'rows': 0}
        with patch('migration.migration_manager.migrate_access', side_effect=fake_migration) as manager:
            self.service.migrate(source, target)
            manager.assert_called_once()
            with self.assertRaises(FileExistsError):
                self.service.migrate(source, target)
            with self.assertRaisesRegex(ValueError, '当前'):
                self.service.migrate(source, self.db)
            self.assertEqual(1, manager.call_count)
        self.assertEqual(original, source.read_bytes())
        self.assertEqual(before, self.state())

    def settings_window(self):
        window = QMainWindow()
        settings = SettingsService(self.settings_path)
        page = SettingsPage(settings, lambda _: None, data_management=self.service)
        window.setCentralWidget(page)
        self.addCleanup(window.close)
        window.show()
        self.app.processEvents()
        controller = DataManagementController(window, page, self.service, lambda _: None, lambda: True)
        return window, page, controller

    def test_settings_buttons_worker_backup_and_export(self):
        window, page, controller = self.settings_window()
        for name in ('accessImport', 'backupDatabase', 'restoreDatabase', 'exportPeople'):
            self.assertTrue(page.findChild(QPushButton, name).isEnabled())
        target = self.root / 'gui-backup.zip'
        with patch.object(QFileDialog, 'getSaveFileName', return_value=(str(target), '')):
            controller.backup()
        self.assertTrue(target.is_file())
        self.assertEqual('', page.error_label.text())
        target = self.root / 'gui-people.xlsx'
        with patch.object(QFileDialog, 'getSaveFileName', return_value=(str(target), '')):
            controller.export_people()
        self.assertTrue(target.is_file())
        self.assertFalse(window._data_operation_running)
        with self.assertRaisesRegex(ValueError, 'worker failure'):
            run_operation(window, lambda: (_ for _ in ()).throw(ValueError('worker failure')), 'test')

    def test_incomplete_rollback_preserves_safety_and_recovery_files(self):
        archive = self.create()
        self.change_live_data()
        original = os.replace
        def fail_replace_and_rollback(source, target):
            if Path(target) == self.images and Path(source).name == 'new':
                raise OSError('Images installation failed')
            if Path(target) == self.db and Path(source).name == 'old':
                raise OSError('database rollback failed')
            return original(source, target)
        with patch('services.backup.os.replace', side_effect=fail_replace_and_rollback):
            with self.assertRaisesRegex(RestoreRecoveryRequiredError, '保留目录'):
                self.backups.restore_backup(archive)
        self.assertTrue(list((self.root / 'Backups').glob('*.zip')))
        self.assertTrue(any(path.joinpath('old').is_file() for path in self.root.glob('.TenderSystem-restore-*')))

    def test_restore_gui_no_is_zero_changes_and_yes_exits_without_settings_write(self):
        archive = self.create()
        self.change_live_data()
        before = self.state()
        window, page, controller = self.settings_window()
        with patch.object(QFileDialog, 'getOpenFileName', return_value=(str(archive), '')), \
             patch.object(QMessageBox, 'question', return_value=QMessageBox.No):
            controller.restore()
        self.assertEqual(before, self.state())
        with patch.object(QFileDialog, 'getOpenFileName', return_value=(str(archive), '')), \
             patch.object(QMessageBox, 'question', return_value=QMessageBox.Yes), \
             patch.object(QMessageBox, 'information') as message, \
             patch.object(window, 'close', wraps=window.close) as close, \
             patch.object(page.settings_service, 'save', wraps=page.settings_service.save) as save:
            controller.restore()
            message.assert_called_once()
            close.assert_called_once()
            save.assert_not_called()
        self.assertTrue(window._restore_completed)
        self.assertEqual('{"dpi":300,"future":"保留"}', self.settings_path.read_text(encoding='utf-8'))

    def test_migration_dialog_confirmation_and_result_without_database_switch(self):
        window, _, _ = self.settings_window()
        dialog = AccessMigrationDialog(window, self.service)
        self.addCleanup(dialog.close)
        source = self.root / 'source.accdb'
        source.write_bytes(b'Access')
        target = self.root / 'migration.db'
        dialog.source.setText(str(source))
        dialog.target.setText(str(target))
        with patch.object(QMessageBox, 'question', return_value=QMessageBox.No), \
             patch.object(self.service, 'migrate') as migrate:
            dialog.start()
            migrate.assert_not_called()
        with patch.object(QMessageBox, 'question', return_value=QMessageBox.Yes), \
             patch.object(self.service, 'migrate', return_value={'rows': 1502}) as migrate:
            dialog.start()
            migrate.assert_called_once_with(source, target)
        self.assertIn('迁移成功', dialog.result_label.text())
        self.assertIn('当前数据库未切换', dialog.result_label.text())
        dialog.target.setText(str(self.db))
        with patch.object(self.service, 'migrate') as migrate:
            dialog.start()
            migrate.assert_not_called()
