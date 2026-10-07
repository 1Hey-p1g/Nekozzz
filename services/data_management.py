"""GUI-facing Access migration and people-only XLSX export, without synchronization."""

from pathlib import Path
import hashlib
import os
import tempfile
import sqlite3
from contextlib import closing

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config.paths import DATA_ROOT, EXCEL_FILE
from database.initializer import check_new_target
from database.runtime import runtime_connection, validate_runtime_path, runtime_maintenance
from database.initializer import publish_staging
from migration.migration_validator import validate_target_structure
from migration.migration_manager import migrate_access
from services.backup import BackupService, publish_file, RestoreRecoveryRequiredError


class DataManagementService:
    def __init__(self, database_path, *, data_root=DATA_ROOT, settings_path=None, migration_resource=EXCEL_FILE):
        self.database_path = validate_runtime_path(database_path)
        self.data_root = Path(data_root).resolve()
        self.migration_resource = Path(migration_resource)
        self.backups = BackupService(database_path, data_root=data_root, settings_path=settings_path)

    def check_migration(self, source_path, target_path):
        source = Path(source_path).expanduser().resolve()
        target = Path(target_path).expanduser().resolve()
        if target == self.database_path:
            raise ValueError('迁移目标不能是当前正在使用的数据库')
        target = check_new_target(target)
        if not source.is_file():
            raise FileNotFoundError(f'Access 源文件不存在：{source}')
        if not self.migration_resource.is_file():
            raise FileNotFoundError(f'程序缺少迁移结构资源：{self.migration_resource}')
        return source, target

    def migrate(self, source_path, target_path):
        source, target = self.check_migration(source_path, target_path)
        return migrate_access(source, target, excel_file=self.migration_resource)

    def _empty_fingerprint(self, path=None):
        path = Path(path or self.database_path)
        if path.is_symlink():
            raise ValueError('默认数据库不能是链接')
        if any(Path(str(path)+suffix).exists() for suffix in ('-wal','-shm','-journal')):
            raise ValueError('默认数据库存在日志文件，不能作为首次初始化空库替换')
        with runtime_connection(path) as connection:
            validate_target_structure(connection, require_empty=True)
        stat = path.stat()
        return (stat.st_ino, stat.st_mtime_ns, hashlib.sha256(path.read_bytes()).hexdigest())

    def can_migrate_as_current(self):
        if self.database_path != self.data_root / 'database' / 'TenderSystem-v1.db':
            return False
        try:
            with runtime_maintenance():
                self._empty_fingerprint()
            return True
        except (OSError, ValueError, sqlite3.Error):
            return False

    def migrate_as_current(self, source_path):
        """Migrate to an independent new file; replace only the unchanged default empty DB."""
        with runtime_maintenance():
            if not self.can_migrate_as_current():
                raise ValueError('仅当前 portable 默认 V1 空数据库允许迁移后替换')
            before = self._empty_fingerprint()
        # Staging is on the same volume. The migration engine publishes only a
        # previously nonexistent file, never INSERTs into the running database.
        work = Path(tempfile.mkdtemp(prefix='.TenderSystem-first-migration-', dir=self.database_path.parent))
        new, old = work / 'migrated.db', work / 'original-empty.db'
        keep = False
        try:
            result = self.migrate(source_path, new)
            with runtime_connection(new) as connection:
                validate_target_structure(connection)
            with runtime_maintenance():
                if self._empty_fingerprint() != before:
                    raise ValueError('当前数据库在迁移期间发生变化，已取消替换')
                # An exclusive SQLite lock detects outstanding readers/writers.
                with closing(sqlite3.connect(self.database_path, timeout=0)) as connection:
                    connection.execute('PRAGMA foreign_keys=ON')
                    connection.execute('BEGIN EXCLUSIVE')
                    validate_target_structure(connection, require_empty=True)
                    connection.rollback()
                if self._empty_fingerprint() != before:
                    raise ValueError('当前数据库在替换前发生变化，已取消替换')
                os.rename(self.database_path, old)
                try:
                    if any(Path(str(self.database_path)+suffix).exists() for suffix in ('-wal','-shm','-journal')):
                        raise ValueError('默认数据库日志在替换瞬间发生变化，已取消替换')
                    # Verify the actual file removed from the live pathname,
                    # closing the check/rename gap. All connections are closed.
                    if self._empty_fingerprint(old) != before:
                        raise ValueError('默认数据库在替换瞬间发生变化，已取消替换')
                    publish_staging(new, self.database_path)
                except BaseException:
                    try:
                        publish_staging(old, self.database_path)
                    except OSError as rollback_error:
                        keep = True
                        raise RestoreRecoveryRequiredError(
                            f'空库迁移替换未完成；原数据库保留在 {old}，请停止使用并恢复：{rollback_error}') from rollback_error
                    raise
            return result
        finally:
            if not keep:
                import shutil
                shutil.rmtree(work, ignore_errors=True)

    def export_people(self, target_path):
        with runtime_connection(self.database_path) as connection:
            columns = [(row['name'], row['type']) for row in connection.execute('PRAGMA table_info(人员表)')]
            # Only this real V1 table; no certificate/project joins or derived columns.
            rows = [tuple(row) for row in connection.execute('SELECT * FROM 人员表 ORDER BY 人员ID')]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '人员表'
        headers = [('人员ID（系统字段，请勿修改）' if name == '人员ID' else
                    '入职日期（进入系统日期）' if name == '进入系统日期' else name) for name, _ in columns]
        sheet.append(headers)
        for index, (name, kind) in enumerate(columns, 1):
            width = max(len(headers[index - 1]) * 1.4, 14)
            width = min(max(width, max((len(str(row[index - 1])) for row in rows if row[index - 1] is not None), default=0)), 42)
            sheet.column_dimensions[get_column_letter(index)].width = width
            cell = sheet.cell(1, index)
            cell.font = Font(bold=True, color='23405D')
            cell.fill = PatternFill('solid', fgColor='E9EFF6')
            cell.alignment = Alignment(vertical='center')
        for row_number, values in enumerate(rows, 2):
            for column_number, (value, (_name, kind)) in enumerate(zip(values, columns), 1):
                cell = sheet.cell(row_number, column_number, value)
                if kind == 'TEXT' and value is not None:
                    # Preserve long numbers and leading zeros. Also keep leading
                    # '=' literal, rather than creating an executable Excel formula.
                    cell.data_type = 's'
                    cell.number_format = '@'
                cell.alignment = Alignment(vertical='top')
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 25
        try:
            path = publish_file(target_path, workbook.save)
        finally:
            workbook.close()
        return path, len(rows)
